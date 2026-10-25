# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).

import base64
import io
from datetime import date

import openpyxl

from odoo.exceptions import UserError
from odoo.tests import tagged
from odoo.tests.common import TransactionCase


@tagged("-at_install", "post_install")
class TestMappingWizard(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.fse = cls.env["financial.statement.eu"]
        cls.bank = cls.env["account.account"].create(
            {"code": "999901", "name": "Test bank", "account_type": "asset_cash"}
        )
        cls.sales = cls.env["account.account"].create(
            {"code": "999902", "name": "Test sales", "account_type": "income"}
        )
        journal = cls.env["account.journal"].search(
            [("company_id", "=", cls.env.company.id), ("type", "=", "general")],
            limit=1,
        )
        move = cls.env["account.move"].create(
            {
                "journal_id": journal.id,
                "date": date(2023, 5, 1),
                "line_ids": [
                    (0, 0, {"account_id": cls.bank.id, "debit": 100}),
                    (0, 0, {"account_id": cls.sales.id, "credit": 100}),
                ],
            }
        )
        move.action_post()

    def _wizard(self):
        return self.env["financial.statement.eu.mapping.wizard"].create(
            {"date_from": date(2023, 1, 1), "date_to": date(2023, 12, 31)}
        )

    def _line(self, code):
        return self.fse.search([("code", "=", code)])

    def _sheet_rows(self, wizard, sheet):
        workbook = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(wizard.export_file))
        )
        return [list(row) for row in workbook[sheet].iter_rows(values_only=True)]

    def _import_csv(self, wizard, text):
        wizard.write(
            {
                "import_file": base64.b64encode(text.encode()),
                "import_filename": "mapping.csv",
            }
        )
        wizard.action_import()
        return {line.account_code: line for line in wizard.line_ids}

    def test_export(self):
        self.sales.financial_statement_eu_debit_id = self._line("E.A1")
        wizard = self._wizard()
        self.assertIn("PP.A8", wizard.instructions)
        wizard.action_export()
        sheets = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(wizard.export_file))
        ).sheetnames
        self.assertEqual(sheets, ["Instructions", "Lines", "Accounts"])
        lines = {row[0]: row for row in self._sheet_rows(wizard, "Lines")[1:]}
        self.assertEqual(lines["E.A1"][2], "yes")
        self.assertEqual(lines["E.A"][2], "no")
        self.assertEqual(lines["E=B"][2], "no")
        self.assertEqual(lines["PP=A9"][2], "yes")
        accounts = {row[0]: row for row in self._sheet_rows(wizard, "Accounts")[1:]}
        self.assertEqual(accounts["999901"][3:8], [100, 0, 100, None, None])
        self.assertEqual(accounts["999902"][5:7], [-100, "E.A1"])

    def test_import(self):
        wizard = self._wizard()
        changes = self._import_csv(
            wizard,
            "Code;Debit line;Credit line;Notes\n"
            "999901;PA.C41;PP.D3a;bank\n"
            "999902;E.A1;;\n",
        )
        self.assertFalse(wizard.error_count)
        self.assertEqual(changes["999901"].notes, "bank")
        wizard.action_apply()
        self.assertEqual(
            self.bank.financial_statement_eu_debit_id, self._line("PA.C41")
        )
        self.assertEqual(
            self.bank.financial_statement_eu_credit_id, self._line("PP.D3a")
        )
        self.assertEqual(self.sales.financial_statement_eu_debit_id, self._line("E.A1"))
        # The same file again changes nothing
        wizard = self._wizard()
        changes = self._import_csv(
            wizard, "Code,Debit line,Credit line\n999902,E.A1,\n"
        )
        self.assertFalse(changes)

    def test_import_errors(self):
        wizard = self._wizard()
        changes = self._import_csv(
            wizard,
            "Code;Debit line;Credit line\n"
            "999901;PA.C4;\n"
            "999902;;\n"
            "000000;E.A1;\n"
            "999903;XX;\n",
        )
        self.assertEqual(wizard.error_count, 4)
        self.assertIn("computed", changes["999901"].error)
        self.assertIn("balance", changes["999902"].error)
        self.assertIn("unknown", changes["000000"].error)
        self.assertIn("not a line", changes["999903"].error)
        with self.assertRaises(UserError):
            wizard.action_apply()
        with self.assertRaises(UserError):
            self._import_csv(self._wizard(), "Code;Line\n999901;E.A1\n")

    def test_round_trip_xlsx(self):
        """The exported file, filled in, imports back."""
        wizard = self._wizard()
        wizard.action_export()
        workbook = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(wizard.export_file))
        )
        sheet = workbook["Accounts"]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value in ("999901", "999902"):
                row[6].value = "PA.C41" if row[0].value == "999901" else "E.A1"
        output = io.BytesIO()
        workbook.save(output)
        wizard.write(
            {
                "import_file": base64.b64encode(output.getvalue()),
                "import_filename": "mapping.xlsx",
            }
        )
        wizard.action_import()
        self.assertFalse(wizard.error_count)
        self.assertEqual(
            set(wizard.line_ids.mapped("account_code")) & {"999901", "999902"},
            {"999901", "999902"},
        )
