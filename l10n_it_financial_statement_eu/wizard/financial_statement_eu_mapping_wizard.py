# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).

import base64
import csv
import io
from datetime import date

import xlsxwriter

from odoo import api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

# Column keys of the "Accounts" sheet, in order; titles in _get_columns()
ACCOUNT_COLUMNS = [
    "code",
    "name",
    "type",
    "debit",
    "credit",
    "balance",
    "debit_line",
    "credit_line",
    "notes",
]


class FinancialStatementEUMappingLine(models.TransientModel):
    _name = "financial.statement.eu.mapping.line"
    _description = "Change of the link of an account to the Financial statement EU"
    _order = "account_code"

    wizard_id = fields.Many2one(
        "financial.statement.eu.mapping.wizard", required=True, ondelete="cascade"
    )
    account_id = fields.Many2one("account.account", readonly=True)
    account_code = fields.Char(string="Code", readonly=True)
    old_debit_id = fields.Many2one(
        "financial.statement.eu", string="Current debit line", readonly=True
    )
    old_credit_id = fields.Many2one(
        "financial.statement.eu", string="Current credit line", readonly=True
    )
    debit_id = fields.Many2one(
        "financial.statement.eu", string="New debit line", readonly=True
    )
    credit_id = fields.Many2one(
        "financial.statement.eu", string="New credit line", readonly=True
    )
    notes = fields.Char(readonly=True)
    error = fields.Char(readonly=True)


class FinancialStatementEUMappingWizard(models.TransientModel):
    """Export the link of the accounts to the lines of the financial
    statements, with instructions for whoever fills it in, a person or an AI
    assistant, and import it back after a check of the changes."""

    _name = "financial.statement.eu.mapping.wizard"
    _description = "Link of the accounts to the Financial statement EU"

    date_from = fields.Date(
        string="From date",
        required=True,
        default=lambda self: date(date.today().year - 1, 1, 1),
    )
    date_to = fields.Date(
        string="To date",
        required=True,
        default=lambda self: date(date.today().year - 1, 12, 31),
    )
    state = fields.Selection(
        [("start", "Start"), ("preview", "Preview")], default="start"
    )
    instructions = fields.Text(compute="_compute_instructions")
    export_file = fields.Binary(readonly=True, attachment=False)
    export_filename = fields.Char(readonly=True)
    import_file = fields.Binary(string="File to import", attachment=False)
    import_filename = fields.Char()
    line_ids = fields.One2many(
        "financial.statement.eu.mapping.line", "wizard_id", string="Changes"
    )
    error_count = fields.Integer(compute="_compute_error_count")

    def _get_columns(self):
        return {
            "code": self.env._("Code"),
            "name": self.env._("Account"),
            "type": self.env._("Type"),
            "debit": self.env._("Debit"),
            "credit": self.env._("Credit"),
            "balance": self.env._("Balance"),
            "debit_line": self.env._("Debit line"),
            "credit_line": self.env._("Credit line"),
            "notes": self.env._("Notes"),
        }

    @api.depends("date_from", "date_to")
    def _compute_instructions(self):
        for wizard in self:
            wizard.instructions = wizard._get_instructions()

    def _get_instructions(self):
        cols = self._get_columns()
        return "\n\n".join(
            [
                self.env._(
                    "How to link the chart of accounts to the financial statements"
                ),
                self.env._(
                    'The file has three sheets. "Lines" lists the lines of the '
                    "financial statements (art. 2424-2425 c.c.): only the lines "
                    'with "yes" in the "Assignable" column can receive accounts, '
                    'the others are totals or are computed. "Accounts" lists the '
                    "accounts of the chart of accounts, with their movements from "
                    "%(date_from)s to %(date_to)s and their current lines. Fill "
                    'in the "%(debit_line)s" and "%(credit_line)s" columns of '
                    '"Accounts" with the codes of the lines, then import the file '
                    "with the same wizard: the changes are shown before they are "
                    "applied.",
                    date_from=self.date_from,
                    date_to=self.date_to,
                    debit_line=cols["debit_line"],
                    credit_line=cols["credit_line"],
                ),
                self.env._(
                    "1. Every account with a balance must have at least one line."
                ),
                self.env._(
                    "2. The debit line receives the balance when it is a debit, "
                    "the credit line when it is a credit. With only one of the "
                    "two, the balance goes there whatever its sign. Fill in both "
                    "only for the accounts that change nature with the sign of "
                    "their balance, such as bank accounts (PA.C41 Bank and postal "
                    "deposits when in debit, PP.D3a Payables to banks when in "
                    "credit) or tax accounts (receivables or payables)."
                ),
                self.env._(
                    "3. Balance sheet accounts go to the lines starting with PA "
                    "(assets) or PP (equity and liabilities), income statement "
                    "accounts to the lines starting with E."
                ),
                self.env._(
                    "4. Accumulated depreciation and allowances go to the line of "
                    "the asset they adjust: the financial statements show assets "
                    "net of them."
                ),
                self.env._(
                    "5. Receivables and payables have separate lines for the "
                    "amounts due within and beyond the next financial year: "
                    "choose by the due date of the account."
                ),
                self.env._(
                    "6. The result of the year is computed from the income "
                    "statement: the accounts that receive it at closing go to "
                    "PP=A9, which does not add them again. The results of "
                    "previous years not yet allocated go to PP.A8 Retained "
                    "earnings."
                ),
                self.env._(
                    '7. Do not change the "%(code)s" column and do not rename '
                    "columns: the import finds the accounts by code and the "
                    "columns by their title. Accounts left out of the file keep "
                    "their lines, an empty cell removes the line. The other "
                    "columns are information only.",
                    code=cols["code"],
                ),
                self.env._(
                    '8. Write the reason of any doubtful choice in the "%(notes)s" '
                    "column: the import shows it next to the change.",
                    notes=cols["notes"],
                ),
                self.env._(
                    "The file to return is the same file, as XLSX, or a CSV file "
                    'with the columns "%(code)s", "%(debit_line)s", '
                    '"%(credit_line)s" and, if useful, "%(notes)s". Whoever fills '
                    "it in, a person or an AI assistant, returns all the accounts "
                    'and uses only the codes of the "Lines" sheet with "yes" in '
                    '"Assignable".',
                    **cols,
                ),
            ]
        )

    def _get_assignable_lines(self):
        """The lines accounts can be linked to, as in the account form."""
        return self.env["financial.statement.eu"].search(
            [
                ("child_ids", "=", False),
                "|",
                ("code", "=", "PP=A9"),
                ("code", "not like", "%=%"),
            ]
        )

    def _get_accounts(self):
        return self.env["account.account"].search(
            [("company_ids", "in", self.env.company.id)], order="code"
        )

    def _get_balances(self):
        """Debit and credit of every account in the period, as the financial
        statements compute them: posted moves, closing moves left out."""
        domain = [
            ("company_id", "=", self.env.company.id),
            ("date", ">=", self.date_from),
            ("date", "<=", self.date_to),
            ("move_id.state", "=", "posted"),
        ]
        if self.env["account.move"].fields_get(allfields=["closing_type"]):
            domain.append(
                ("move_id.closing_type", "not in", ("closing", "loss_profit"))
            )
        groups = self.env["account.move.line"]._read_group(
            domain, ["account_id"], ["debit:sum", "credit:sum"]
        )
        return {account.id: (debit, credit) for account, debit, credit in groups}

    def action_export(self):
        self.ensure_one()
        cols = self._get_columns()
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        bold = workbook.add_format({"bold": True})
        wrap = workbook.add_format({"text_wrap": True, "valign": "top"})
        amount = workbook.add_format({"num_format": "#,##0.00"})

        sheet = workbook.add_worksheet(self.env._("Instructions"))
        sheet.set_column(0, 0, 120, wrap)
        for row, paragraph in enumerate(self.instructions.split("\n\n")):
            sheet.write(row, 0, paragraph, bold if row == 0 else wrap)

        sheet = workbook.add_worksheet(self.env._("Lines"))
        sheet.write_row(
            0, 0, [cols["code"], self.env._("Line"), self.env._("Assignable")], bold
        )
        sheet.set_column(0, 0, 10)
        sheet.set_column(1, 1, 100)
        sheet.set_column(2, 2, 12)
        assignable = self._get_assignable_lines()
        lines = self.env["financial.statement.eu"].search([("ordinary", "=", True)])
        for row, line in enumerate(lines, start=1):
            path = line.get_parent_path()[: -len(" / ")]
            yes = line in assignable
            sheet.write_row(
                row,
                0,
                [line.code, path, self.env._("yes") if yes else self.env._("no")],
            )

        sheet = workbook.add_worksheet(self.env._("Accounts"))
        sheet.write_row(0, 0, [cols[key] for key in ACCOUNT_COLUMNS], bold)
        sheet.set_column(0, 0, 14)
        sheet.set_column(1, 1, 50)
        sheet.set_column(2, 2, 22)
        sheet.set_column(3, 5, 14, amount)
        sheet.set_column(6, 7, 12)
        sheet.set_column(8, 8, 60)
        types = dict(
            self.env["account.account"]
            ._fields["account_type"]
            ._description_selection(self.env)
        )
        balances = self._get_balances()
        for row, account in enumerate(self._get_accounts(), start=1):
            debit, credit = balances.get(account.id, (0.0, 0.0))
            sheet.write_row(
                row,
                0,
                [
                    account.code,
                    account.name,
                    types.get(account.account_type, ""),
                    debit,
                    credit,
                    debit - credit,
                    account.financial_statement_eu_debit_id.code or "",
                    account.financial_statement_eu_credit_id.code or "",
                    "",
                ],
            )
        workbook.close()
        self.write(
            {
                "export_file": base64.b64encode(output.getvalue()),
                "export_filename": self.env._("financial_statement_eu_mapping.xlsx"),
            }
        )
        return self._reopen()

    def _reopen(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    @staticmethod
    def _cell(value):
        """Codes as text, also when a spreadsheet turned them into numbers."""
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()

    def _read_rows(self):
        """The rows of the file as lists of text cells, header included: the
        "Accounts" sheet of an XLSX file, or a CSV file."""
        data = base64.b64decode(self.import_file)
        if (self.import_filename or "").lower().endswith(".csv"):
            text = data.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return [
                [self._cell(c) for c in row]
                for row in csv.reader(io.StringIO(text), dialect)
            ]
        if openpyxl is None:  # pragma: no cover
            raise UserError(self.env._("Reading XLSX files needs openpyxl."))
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True
            )
        except Exception as error:
            raise UserError(
                self.env._("The file is neither an XLSX nor a CSV file.")
            ) from error
        # The sheet with the line columns: "Lines" has a code column too
        titles = {self._get_columns()["debit_line"].lower(), "debit line"}
        for sheet in workbook.worksheets:
            rows = [
                [self._cell(c) for c in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            if rows and titles & {c.lower() for c in rows[0]}:
                return rows
        return []

    def _find_columns(self, header):
        """Position of each column, by its title in the user language or in
        English."""
        cols = self._get_columns()
        english = {key: key.replace("_", " ") for key in ACCOUNT_COLUMNS}
        english["name"] = "account"
        header = [c.lower() for c in header]
        found = {}
        for key in ACCOUNT_COLUMNS:
            for title in (cols[key].lower(), english[key]):
                if title in header:
                    found[key] = header.index(title)
                    break
        missing = [
            cols[k] for k in ("code", "debit_line", "credit_line") if k not in found
        ]
        if missing:
            raise UserError(
                self.env._(
                    "The file has no column %(columns)s: the titles must stay "
                    "as in the exported file.",
                    columns=", ".join(missing),
                )
            )
        return found

    def action_import(self):
        self.ensure_one()
        if not self.import_file:
            raise UserError(self.env._("Choose the file to import."))
        rows = self._read_rows()
        if not rows:
            raise UserError(self.env._("The file has no sheet with the accounts."))
        pos = self._find_columns(rows[0])
        accounts = {a.code: a for a in self._get_accounts()}
        assignable = {line.code: line for line in self._get_assignable_lines()}
        all_lines = {
            line.code: line for line in self.env["financial.statement.eu"].search([])
        }
        balances = self._get_balances()
        fse = self.env["financial.statement.eu"]
        values = []
        seen = set()
        for row in rows[1:]:
            row = row + [""] * (len(ACCOUNT_COLUMNS) - len(row))
            code = row[pos["code"]]
            if not code:
                continue
            notes = row[pos["notes"]] if "notes" in pos else ""
            account = accounts.get(code)
            errors = []
            if not account:
                errors.append(self.env._("unknown account"))
            elif code in seen:
                errors.append(self.env._("account repeated in the file"))
            seen.add(code)
            new = {}
            for key in ("debit_line", "credit_line"):
                line_code = row[pos[key]]
                new[key] = assignable.get(line_code, fse)
                if line_code and line_code not in assignable:
                    errors.append(
                        self.env._("%(code)s is not a line", code=line_code)
                        if line_code not in all_lines
                        else self.env._(
                            "%(code)s is a total or is computed", code=line_code
                        )
                    )
            if account and not (new["debit_line"] or new["credit_line"]):
                debit, credit = balances.get(account.id, (0.0, 0.0))
                if self.env.company.currency_id.compare_amounts(debit, credit):
                    errors.append(self.env._("the account has a balance and no line"))
            old_debit = account.financial_statement_eu_debit_id if account else fse
            old_credit = account.financial_statement_eu_credit_id if account else fse
            changed = (new["debit_line"], new["credit_line"]) != (old_debit, old_credit)
            if errors or changed or notes:
                values.append(
                    {
                        "wizard_id": self.id,
                        "account_id": account.id if account else False,
                        "account_code": code,
                        "old_debit_id": old_debit.id,
                        "old_credit_id": old_credit.id,
                        "debit_id": new["debit_line"].id,
                        "credit_id": new["credit_line"].id,
                        "notes": notes,
                        "error": ", ".join(errors),
                    }
                )
        self.line_ids.unlink()
        self.env["financial.statement.eu.mapping.line"].create(values)
        self.state = "preview"
        return self._reopen()

    @api.depends("line_ids.error")
    def _compute_error_count(self):
        for wizard in self:
            wizard.error_count = len(wizard.line_ids.filtered("error"))

    def action_apply(self):
        self.ensure_one()
        if self.error_count:
            raise UserError(
                self.env._("Correct the errors in the file and import it again.")
            )
        for line in self.line_ids.filtered("account_id"):
            line.account_id.write(
                {
                    "financial_statement_eu_debit_id": line.debit_id.id,
                    "financial_statement_eu_credit_id": line.credit_id.id,
                }
            )
        return {"type": "ir.actions.act_window_close"}

    def action_back(self):
        self.ensure_one()
        self.line_ids.unlink()
        self.write({"state": "start", "import_file": False, "import_filename": False})
        return self._reopen()
